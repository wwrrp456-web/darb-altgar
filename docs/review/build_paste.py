# ملف وورد واحد بنصوص عادية (بلا جداول) للنسخ واللصق في محادثة المراجع: python3 docs/review/build_texts_review.py json ثم python3 docs/review/build_paste.py
import json
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

B = json.load(open('/tmp/batches.json', encoding='utf-8'))
FONT = 'Tahoma'
doc = Document(); sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7); sec.left_margin = sec.right_margin = Cm(2); sec.top_margin = sec.bottom_margin = Cm(2)
st = doc.styles['Normal']; st.font.name = FONT; st.font.size = Pt(12)
for a in ('cs', 'ascii', 'hAnsi'): st.element.rPr.rFonts.set(qn('w:' + a), FONT)


def para(text='', bold=False, size=None, color=None, after=6):
    p = doc.add_paragraph(); p._p.get_or_add_pPr().append(OxmlElement('w:bidi')); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(after)
    if text:
        r = p.add_run(text); r.bold = bold; r.font.name = FONT; r._element.get_or_add_rPr().rFonts.set(qn('w:cs'), FONT); r._element.rPr.append(OxmlElement('w:rtl'))
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = color
    return p


NAVY, GOLD = RGBColor(0x2F, 0x4A, 0x6D), RGBColor(0xB9, 0x8A, 0x43)
RULES = '''أنت مراجع لغوي وأدبي للعبة تعليمية عربية اسمها «درب التجّار»، يتعلم فيها اللاعب (ناشئ أو شاب) أصول التجارة بتوجيه شخصية «عمّ سالم»، كاتب حسابات جدّه، ويخاطب اللاعب بـ«يا ولدي». سأرسل لك النصوص على دفعات صغيرة. مطلوب منك في كل دفعة:
١) تدقيق لغوي ونحوي وإملائي.
٢) إعادة صياغة أدبية بفصحى مبسّطة يفهمها ابن ١٢ سنة وأبوه، بلسان دافئ أبوي واحد لعمّ سالم، بلا تكلّف ولا وعظ ثقيل.
٣) جعل النصوص مترابطة تخدم اللاعب: درس واحد في كل لحظة؛ كل شخصية تُعرَّف بجملة عند أول ظهور أو تُستبدل بعبارة عامة؛ لا إشارة إلى شيء قبل أن يُروى؛ لا خروج عن عالم الحكاية (لا «في لعبتنا» داخلها)؛ وكل نص يخدم قراراً أو فهماً عند اللاعب، فإن لم يخدمه فاقترح حذفه وسبّب ذلك.
قواعد ثابتة:
• كل نص له رمز بين قوسين مربعين مثل [L1.intro]. لا تغيّر الرمز.
• الرموز بين أقواس معقوفة مثل {جار} و{م} و{boss}، والأرقام والنسب والمبالغ: تبقى كما هي حرفياً (تتبدل داخل اللعبة بحسب البلد).
• النصوص المعلَّمة (مقفل للمختص) مضمونها شرعي أو تاريخي: دقّق لغتها فقط، واكتب أي تحسين أسلوبي مقترحاً دون تغيير المعنى، ولا تضف حكماً ولا حديثاً ولا معلومة.
• لا تضف أحكاماً دينية أو معلومات تاريخية من عندك؛ وإن نقصتك معلومة فاسأل.
• تعبيرات حساسة (مثل وصف إنسان بأنه «عبد» أو «روح العبد») اقترح لها بديلاً يحفظ الفكرة دون تحقير.
• هذه أسماء مشتركة بين الحكايات ولا تتغير: سعيد، الحاج عمران، ظافر، حمدان، مبارك، عامر، سهيل، غانم، حمود، والحاج صالح الصيرفي (دائن الجدّ). وكذلك «المقرض الحكيم صالح» يلتبس بالحاج صالح فاقترح له اسماً مختلفاً.
صيغة ردّك لكل نص يحتاج تعديلاً:
[الرمز] النص المقترح
وما كان سليماً فلا تكتبه. وما ترى حذفه فاكتب: [الرمز] حذف — والسبب. وإن كانت لك ملاحظة ترابط فاكتبها في آخر الرد تحت عنوان «ملاحظات».
كل نص يسبقه سطر «السياق:» يبيّن من يتكلم ومتى يظهر؛ هو للتوضيح فقط فلا تُعِده في ردّك.\nأجبني بكلمة «جاهز» لأرسل لك الدفعة الأولى.'''
para('درب التجّار — نصوص للنسخ واللصق عند المراجع', True, 18, NAVY, 10)
para('الصفحة الأولى: رسالة التعليمات، تُنسخ وتُرسل مرة واحدة في بداية المحادثة. بعدها كل دفعة في صفحة مستقلة، تُنسخ بين سطري «ابدأ الدفعة» و«انتهت الدفعة» وتُرسل. أعطِ المراجع الدفعة التالية بعد أن يُنهي السابقة، وأعدها إليّ بردوده لأطبقها.', False, 10, None, 12)
para('═══ رسالة التعليمات (انسخ من هنا) ═══', True, 11, GOLD, 8)
for ln in RULES.split('\n'): para(ln)
para('═══ انتهت التعليمات ═══', True, 11, GOLD, 8)
for i, b in enumerate(B, 1):
    doc.add_page_break()
    para('═══ ابدأ الدفعة %d من %d: %s ═══' % (i, len(B), b['title']), True, 12, GOLD, 10)
    for r in b['rows']:
        tag = '[%s]' % r['id']
        if r['lock']: tag += ' (مقفل للمختص)'
        para(tag + ' السياق: ' + r['ctx'] + (' (مقفل للمختص)' if False else ''), False, 9, RGBColor(0x62, 0x6D, 0x74), 1)
        para(r['text'].replace('\n', ' '), after=8)
    para('═══ انتهت الدفعة %d ═══' % i, True, 12, GOLD, 4)
doc.save('docs/review/TEXTS_PASTE.docx'); print('ok')
